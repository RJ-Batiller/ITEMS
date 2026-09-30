"""Reusable report-generation services."""

import csv
from io import StringIO
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_csv(headers, rows):
    """Build a UTF-8 CSV document with a BOM for spreadsheet compatibility."""
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    writer.writerows(rows)
    return "\ufeff" + output.getvalue()


def build_xlsx(headers, rows, sheet_title="Report"):
    """Build an Excel workbook with readable widths and a filterable header."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C8102E")
        cell.alignment = Alignment(vertical="top")

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_letter].width = min(max(longest + 2, 12), 55)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
